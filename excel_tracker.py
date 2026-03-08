import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime
import logging
import os

logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')
logger = logging.getLogger(__name__)

class VehiclePriceTracker:
    """Manage Excel workbook for vehicle pricing tracking."""

    def __init__(self, filename='vehicle_pricing_tracker.xlsx'):
        self.filename = filename
        self.workbook = None
        self.worksheet = None

    def create_new_tracker(self):
        """Create a new tracking workbook."""
        self.workbook = openpyxl.Workbook()
        # Remove default sheet
        if 'Sheet' in self.workbook.sheetnames:
            self.workbook.remove(self.workbook['Sheet'])

        # Create a summary/metadata sheet
        self._create_metadata_sheet()
        logger.info(f"Created new tracker: {self.filename}")

    def load_tracker(self):
        """Load existing tracking workbook."""
        if os.path.exists(self.filename):
            self.workbook = openpyxl.load_workbook(self.filename)
            logger.info(f"Loaded existing tracker: {self.filename}")
        else:
            self.create_new_tracker()

    def _create_metadata_sheet(self):
        """Create a metadata sheet with information about the tracker."""
        ws = self.workbook.create_sheet('Metadata', 0)

        ws['A1'] = 'Vehicle Pricing Tracker'
        ws['A1'].font = Font(bold=True, size=14)

        ws['A3'] = 'Description:'
        ws['A4'] = 'Weekly tracking of MSRP, Invoice, and Destination Fee for top 75 US vehicles'

        ws['A6'] = 'Model Years Tracked:'
        ws['A7'] = '2024, 2025, 2026 (or available years)'

        ws['A9'] = 'Data Columns:'
        headers = ['Brand', 'Model', 'Trim', 'Model Year', 'MSRP', 'Invoice', 'Destination Fee', 'Scraped Date']
        for idx, header in enumerate(headers, 1):
            ws.cell(row=10, column=idx, value=header)
            ws.cell(row=10, column=idx).font = Font(bold=True)

        ws.column_dimensions['A'].width = 20

    def add_weekly_sheet(self, data, week_date=None):
        """
        Add a new weekly sheet with pricing data.

        Args:
            data: List of dicts with pricing information
            week_date: Date for the week (defaults to today)
        """
        if week_date is None:
            week_date = datetime.now()

        # Create sheet name based on week
        sheet_name = week_date.strftime('Week_%Y_%m_%d')[:31]  # Excel limit 31 chars

        if sheet_name in self.workbook.sheetnames:
            logger.warning(f"Sheet {sheet_name} already exists, removing old one")
            self.workbook.remove(self.workbook[sheet_name])

        ws = self.workbook.create_sheet(sheet_name)
        self._format_weekly_sheet(ws, data, week_date)

        logger.info(f"Added weekly sheet: {sheet_name}")

    def _format_weekly_sheet(self, ws, data, week_date):
        """Format the weekly sheet with headers and data."""
        # Add title
        ws['A1'] = f"Vehicle Pricing - Week of {week_date.strftime('%B %d, %Y')}"
        ws['A1'].font = Font(bold=True, size=12)

        # Add headers
        headers = ['Rank', 'Brand', 'Model', 'Trim', 'Model Year', 'MSRP', 'Invoice', 'Destination Fee', 'Total Cost', 'Scraped Date']
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_num)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
            cell.alignment = Alignment(horizontal="center", vertical="center")

        # Add data rows
        row_num = 4
        for item in data:
            ws.cell(row=row_num, column=1, value=item.get('rank'))
            ws.cell(row=row_num, column=2, value=item.get('brand'))
            ws.cell(row=row_num, column=3, value=item.get('model'))
            ws.cell(row=row_num, column=4, value=item.get('trim'))
            ws.cell(row=row_num, column=5, value=item.get('model_year'))

            # Price columns - format as currency
            msrp = item.get('msrp')
            invoice = item.get('invoice')
            dest_fee = item.get('destination_fee')

            ws.cell(row=row_num, column=6, value=msrp)
            ws.cell(row=row_num, column=7, value=invoice)
            ws.cell(row=row_num, column=8, value=dest_fee)

            # Total Cost calculation
            if msrp and dest_fee:
                total = msrp + dest_fee
                ws.cell(row=row_num, column=9, value=total)

            ws.cell(row=row_num, column=10, value=item.get('scraped_date'))

            row_num += 1

        # Format columns
        self._format_columns(ws, row_num - 1)

    def _format_columns(self, ws, last_row):
        """Format columns with appropriate widths and number formats."""
        column_widths = {
            'A': 6,   # Rank
            'B': 12,  # Brand
            'C': 15,  # Model
            'D': 20,  # Trim
            'E': 12,  # Model Year
            'F': 12,  # MSRP
            'G': 12,  # Invoice
            'H': 16,  # Destination Fee
            'I': 12,  # Total Cost
            'J': 20   # Scraped Date
        }

        for col, width in column_widths.items():
            ws.column_dimensions[col].width = width

        # Format currency columns (F, G, H, I)
        currency_cols = ['F', 'G', 'H', 'I']
        for col in currency_cols:
            for row in range(4, last_row + 1):
                cell = ws[f'{col}{row}']
                if cell.value is not None:
                    cell.number_format = '$#,##0.00'
                cell.alignment = Alignment(horizontal="right")

        # Center align certain columns
        for row in range(4, last_row + 1):
            ws[f'A{row}'].alignment = Alignment(horizontal="center")
            ws[f'E{row}'].alignment = Alignment(horizontal="center")

    def save(self):
        """Save the workbook."""
        self.workbook.save(self.filename)
        logger.info(f"Saved tracker: {self.filename}")

    def add_comparison_sheet(self, comparisons):
        """
        Add a sheet comparing prices across model years.

        Args:
            comparisons: Dict with vehicle comparison data
        """
        if 'Comparisons' in self.workbook.sheetnames:
            self.workbook.remove(self.workbook['Comparisons'])

        ws = self.workbook.create_sheet('Comparisons')

        ws['A1'] = 'Year-over-Year Pricing Comparison'
        ws['A1'].font = Font(bold=True, size=12)

        # Headers
        headers = ['Brand', 'Model', 'Trim', '2024 MSRP', '2025 MSRP', '2026 MSRP', '2024-2025 Change', '2025-2026 Change']
        for col_num, header in enumerate(headers, 1):
            cell = ws.cell(row=3, column=col_num)
            cell.value = header
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")

        # Add data
        row_num = 4
        for vehicle_key, comp_data in comparisons.items():
            ws.cell(row=row_num, column=1, value=comp_data.get('brand'))
            ws.cell(row=row_num, column=2, value=comp_data.get('model'))
            ws.cell(row=row_num, column=3, value=comp_data.get('trim'))

            # Add year-specific MSRP values
            msrp_2024 = comp_data.get('msrp_2024')
            msrp_2025 = comp_data.get('msrp_2025')
            msrp_2026 = comp_data.get('msrp_2026')

            ws.cell(row=row_num, column=4, value=msrp_2024)
            ws.cell(row=row_num, column=5, value=msrp_2025)
            ws.cell(row=row_num, column=6, value=msrp_2026)

            # Calculate changes
            if msrp_2024 and msrp_2025:
                change_24_25 = msrp_2025 - msrp_2024
                ws.cell(row=row_num, column=7, value=change_24_25)

            if msrp_2025 and msrp_2026:
                change_25_26 = msrp_2026 - msrp_2025
                ws.cell(row=row_num, column=8, value=change_25_26)

            row_num += 1

        # Format columns
        for col in ['D', 'E', 'F', 'G', 'H']:
            for row in range(4, row_num):
                cell = ws[f'{col}{row}']
                if cell.value is not None:
                    cell.number_format = '$#,##0.00'
                cell.alignment = Alignment(horizontal="right")

        # Set column widths
        for col, width in [('A', 12), ('B', 15), ('C', 20), ('D', 12), ('E', 12), ('F', 12), ('G', 14), ('H', 14)]:
            ws.column_dimensions[col].width = width
